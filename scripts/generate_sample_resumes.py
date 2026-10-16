"""
Generate realistic sample resumes in PDF and DOCX formats for testing and demonstration.
"""

from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "sample_resumes"
SAMPLES_DIR.mkdir(exist_ok=True)


def generate_backend_pdf():
    pdf_path = SAMPLES_DIR / "sample_john_doe_backend.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=4
    )
    contact_style = ParagraphStyle(
        'ContactStyle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        spaceAfter=14
    )
    heading_style = ParagraphStyle(
        'HeadingStyle',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#334155"),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyStyle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#334155"),
        spaceAfter=6
    )

    story = [
        Paragraph("John Doe", title_style),
        Paragraph("Email: john.doe@email.com | Phone: +1 (555) 345-6789 | Location: San Francisco, CA<br/>LinkedIn: linkedin.com/in/johndoe | GitHub: github.com/johndoe", contact_style),
        
        Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style),
        Paragraph("Results-driven Senior Backend Engineer with over 5 years of experience architecting high-throughput microservices and RESTful APIs using Python, Flask, FastAPI, and PostgreSQL. Experienced in distributed systems, Docker containerization, and AWS cloud infrastructure.", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>CORE COMPETENCIES & TECHNICAL SKILLS</b>", heading_style),
        Paragraph("<b>Programming Languages:</b> Python, SQL, Bash, Go<br/>"
                  "<b>Frameworks & Web:</b> Flask, FastAPI, Django, REST API, Microservices<br/>"
                  "<b>Databases & Caching:</b> PostgreSQL, Redis, MySQL<br/>"
                  "<b>Cloud & DevOps:</b> AWS, Docker, Kubernetes, CI/CD, Git, Linux<br/>"
                  "<b>Methodologies & Soft Skills:</b> Agile, Scrum, System Design, Problem Solving, Leadership", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>PROFESSIONAL EXPERIENCE</b>", heading_style),
        Paragraph("<b>Software Engineer at TechCorp Solutions</b> (2021 - Present)<br/>"
                  "• Architected core payment microservices handling 2M+ transactions daily using Python, Flask, and PostgreSQL.<br/>"
                  "• Reduced database query latency by 45% through PostgreSQL index tuning and Redis caching.<br/>"
                  "• Configured CI/CD deployment pipelines using Docker and AWS ECS.", body_style),
        Paragraph("<b>Backend Developer at CloudScale Inc</b> (2019 - 2021)<br/>"
                  "• Built RESTful APIs in Flask and automated database migrations with SQLAlchemy.<br/>"
                  "• Containerized 12 internal microservices using Docker and orchestrated deployments on Kubernetes.", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>EDUCATION</b>", heading_style),
        Paragraph("<b>Bachelor of Science in Computer Science</b> (2015 - 2019)<br/>"
                  "University of California, Berkeley<br/>"
                  "GPA: 3.8/4.0", body_style)
    ]

    doc.build(story)
    print(f"Generated PDF: {pdf_path}")


def generate_ai_docx():
    docx_path = SAMPLES_DIR / "sample_sarah_connor_ai.docx"
    doc = docx.Document()

    # Header
    doc.add_heading("Sarah Connor", level=1)
    p_contact = doc.add_paragraph("Email: sarah.connor@ai-labs.org | Phone: +1 (555) 987-6543 | Location: Boston, MA\n"
                                  "LinkedIn: linkedin.com/in/sarah-connor-ai | GitHub: github.com/sarahconnor")
    
    # Summary
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph("Lead AI & Machine Learning Engineer with 4 years of experience delivering production deep learning, Natural Language Processing (NLP), and computer vision pipelines using Python, PyTorch, spaCy, and Hugging Face.")

    # Skills
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(
        "• Programming: Python, R, C++, SQL\n"
        "• Machine Learning & AI: PyTorch, TensorFlow, spaCy, Scikit-learn, Pandas, NumPy, NLP, LLMs, Computer Vision\n"
        "• Cloud & Tools: AWS, Docker, Git, MLflow, Linux\n"
        "• Soft Skills: Team Management, Communication, Critical Thinking"
    )

    # Experience
    doc.add_heading("Work Experience", level=2)
    p_exp1 = doc.add_paragraph()
    p_exp1.add_run("Machine Learning Engineer at DeepMindset Labs\n").bold = True
    p_exp1.add_run("2022 - Present | Boston, MA\n").italic = True
    p_exp1.add_run("• Developed information extraction pipelines using spaCy and Transformers for enterprise document analysis.\n"
                   "• Fine-tuned large language models (LLMs) improving entity extraction accuracy by 28%.\n"
                   "• Deployed PyTorch models as low-latency microservices with FastAPI and Docker.")

    p_exp2 = doc.add_paragraph()
    p_exp2.add_run("Data Scientist at DataSphere Analytics\n").bold = True
    p_exp2.add_run("2020 - 2022 | Cambridge, MA\n").italic = True
    p_exp2.add_run("• Built predictive classification and regression models using Scikit-learn and Pandas.\n"
                   "• Processed 500GB+ datasets using SQL and Apache Spark.")

    # Education
    doc.add_heading("Education", level=2)
    p_edu = doc.add_paragraph()
    p_edu.add_run("Master of Science in Artificial Intelligence\n").bold = True
    p_edu.add_run("Massachusetts Institute of Technology (MIT)\n")
    p_edu.add_run("Graduation Year: 2020 | GPA: 3.9/4.0")

    doc.save(str(docx_path))
    print(f"Generated DOCX: {docx_path}")


def generate_fullstack_pdf():
    pdf_path = SAMPLES_DIR / "sample_alex_kumar_fullstack.pdf"
    doc = SimpleDocTemplate(str(pdf_path), pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=20, leading=24, textColor=colors.HexColor("#0f172a"))
    contact_style = ParagraphStyle('ContactStyle', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#475569"), spaceAfter=14)
    heading_style = ParagraphStyle('HeadingStyle', parent=styles['Heading2'], fontSize=13, leading=16, textColor=colors.HexColor("#334155"), spaceBefore=10, spaceAfter=6)
    body_style = ParagraphStyle('BodyStyle', parent=styles['Normal'], fontSize=10, leading=14, textColor=colors.HexColor("#334155"), spaceAfter=6)

    story = [
        Paragraph("Alex Kumar", title_style),
        Paragraph("Email: alex.kumar@devmail.com | Phone: +1 415 555 0192 | Location: Seattle, WA<br/>LinkedIn: linkedin.com/in/alexkumar | GitHub: github.com/alexkumar-dev", contact_style),
        
        Paragraph("<b>PROFESSIONAL SUMMARY</b>", heading_style),
        Paragraph("Full Stack Developer with 3+ years of experience building responsive, accessible web applications using React, TypeScript, Node.js, Express, and PostgreSQL. Passionate about performant user interfaces and test-driven development.", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>TECHNICAL SKILLS</b>", heading_style),
        Paragraph("<b>Languages:</b> JavaScript, TypeScript, Python, HTML5, CSS3, SQL<br/>"
                  "<b>Frontend:</b> React, Next.js, Redux, Tailwind CSS, Bootstrap<br/>"
                  "<b>Backend & DB:</b> Node.js, Express, PostgreSQL, MongoDB, GraphQL, REST API<br/>"
                  "<b>Testing & DevOps:</b> Jest, Cypress, Git, GitHub Actions, Docker", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>EXPERIENCE</b>", heading_style),
        Paragraph("<b>Full Stack Developer at WebCraft Innovations</b> (2022 - Present)<br/>"
                  "• Designed and launched dynamic dashboard interfaces with React, TypeScript, and Tailwind CSS.<br/>"
                  "• Developed RESTful backend APIs with Node.js, Express, and PostgreSQL.<br/>"
                  "• Wrote automated end-to-end and unit test suites with Jest and Cypress.", body_style),
        Paragraph("<b>Junior Developer at StartupHub</b> (2021 - 2022)<br/>"
                  "• Developed responsive UI components and integrated third-party payment APIs.<br/>"
                  "• Participated in daily Agile standups and code reviews.", body_style),
        
        Spacer(1, 8),
        Paragraph("<b>EDUCATION</b>", heading_style),
        Paragraph("<b>Bachelor of Technology in Computer Engineering</b> (2017 - 2021)<br/>"
                  "University of Washington<br/>"
                  "Grade: 3.7 GPA", body_style)
    ]

    doc.build(story)
    print(f"Generated PDF: {pdf_path}")


if __name__ == "__main__":
    generate_backend_pdf()
    generate_ai_docx()
    generate_fullstack_pdf()
    print("All sample resumes generated successfully!")
